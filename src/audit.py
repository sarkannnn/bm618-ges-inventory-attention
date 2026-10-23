"""
BM618: Az Verili GES Envanterinde Attention, Transformer ve Kanıta Dayalı LLM
Görev 1.1 — Denetlenebilir Veri Tablosu ve Ham Veri Denetim Hattı (Audit Pipeline)
"""

import os
import sys
import logging
from typing import List, Dict, Any, Tuple
import openpyxl
import pandas as pd
import numpy as np

# Ensure UTF-8 output
sys.stdout.reconfigure(encoding='utf-8')

# District typo correction dictionary (audited)
DISTRICT_CORRECTIONS = {
    'Atnavutköy': 'Arnavutköy',
    'Sultangazi ': 'Sultangazi',
}

RAW_FILES = {
    2022: '2022-yl-kurulumu-tamamlanan-gune-enerji-sistemleri.xlsx',
    2023: '2023_kurulumu-tamamlanan-gune-enerji-sistemleri.xlsx',
    2024: '2024-yl-kurulumu-tamamlanan-gune-enerji-sistemleri.xlsx',
    2025: '2025-yl-kurulumu-tamamlanan-gune-enerji-sistemleri.xlsx',
}

# Ground truth values from Yönerge and Öğretim Elemanı Notları for strict automated assertion
EXPECTED_METRICS = {
    2022: {'count': 20, 'sum': 3834.8, 'districts': 15, 'median': 106.0, 'min': 20.0, 'max': 574.2, 'status': {'Aktif': 20}},
    2023: {'count': 11, 'sum': 984.7, 'districts': 9, 'median': 25.0, 'min': 6.6, 'max': 305.0, 'status': {'Aktif': 11}},
    2024: {'count': 46, 'sum': 3887.0, 'districts': 21, 'median': 36.0, 'min': 3.0, 'max': 484.0, 'status': {'Aktif': 45, '(Aktif / Pasif)': 1}},
    2025: {'count': 76, 'sum': 8418.0, 'districts': 34, 'median': 63.0, 'min': 3.0, 'max': 574.0, 'status': {'Aktif': 59, 'Pasif': 15, 'Aktif - Pasif': 2}},
}


def clean_spaces(text: str) -> str:
    """Normalize multiple whitespace characters and strip edges."""
    if text is None:
        return ""
    return " ".join(str(text).strip().split())


def clean_district(raw_district: str) -> Tuple[str, List[str]]:
    """Clean district name and flag corrections."""
    flags = []
    if raw_district is None:
        return "", ["DISTRICT_MISSING"]
    
    cleaned = clean_spaces(raw_district)
    if raw_district in DISTRICT_CORRECTIONS:
        corrected = DISTRICT_CORRECTIONS[raw_district]
        flags.append(f"DISTRICT_TYPO_FIXED:{raw_district}->{corrected}")
        cleaned = corrected
    elif cleaned in DISTRICT_CORRECTIONS:
        corrected = DISTRICT_CORRECTIONS[cleaned]
        flags.append(f"DISTRICT_TYPO_FIXED:{cleaned}->{corrected}")
        cleaned = corrected
        
    return cleaned, flags


def clean_facility_name(raw_name: str) -> Tuple[str, List[str]]:
    """Normalize facility name while preserving stage numbers and identifiers."""
    flags = []
    if raw_name is None:
        return "", ["NAME_MISSING"]
    
    cleaned = clean_spaces(raw_name)
    return cleaned, flags


def clean_status(raw_status: str) -> Tuple[str, List[str]]:
    """
    Standardize status field.
    Öğretim Elemanı Notu:
    - 2024 '(Aktif / Pasif)' bir aktiflik etiketi sayılmamalıdır (Belirsiz).
    - 2025 'Aktif - Pasif' 2 adet karma kayıttır (Karma).
    - 2025 'Pasif' 15 adettir.
    - 'Aktif' 135 adettir.
    """
    flags = []
    if raw_status is None:
        return "Bilinmiyor", ["STATUS_NULL"]
    
    st = clean_spaces(raw_status)
    if st == "Aktif":
        return "Aktif", []
    elif st == "Pasif":
        return "Pasif", []
    elif st == "(Aktif / Pasif)":
        flags.append("STATUS_SUBHEADER_AMBIGUOUS")
        return "Belirsiz", flags
    elif st in ["Aktif - Pasif", "Aktif-Pasif"]:
        flags.append("STATUS_HYBRID_ACTIVE_PASSIVE")
        return "Karma", flags
    else:
        flags.append(f"STATUS_UNKNOWN:{st}")
        return st, flags


def parse_numeric_power(val: Any) -> Tuple[float, bool]:
    """Parse numeric power in kWp. Returns (value, is_valid)."""
    if val is None:
        return None, False
    try:
        s = str(val).strip().replace(',', '.')
        num = float(s)
        return round(num, 4), True
    except (ValueError, TypeError):
        return None, False


def run_audit(raw_dir: str, output_data_dir: str, output_log_dir: str) -> pd.DataFrame:
    """
    Executes Görev 1.1 audit pipeline:
    - Ingests raw Excel files without modifying them.
    - Logs all header, subheader, and invalid rows.
    - Standardizes columns to the required schema:
      record_id, source_file, sheet, excel_row, file_year,
      district_raw, district_clean, name_raw, name_clean,
      power_kwp, status_raw, status_clean, entity_group_id, quality_flag
    - Asserts that extracted data matches official assignment metrics exactly.
    """
    os.makedirs(output_data_dir, exist_ok=True)
    os.makedirs(output_log_dir, exist_ok=True)
    
    log_file = os.path.join(output_log_dir, "data_audit.log")
    logger = logging.getLogger("DataAudit")
    logger.setLevel(logging.INFO)
    logger.handlers.clear()
    
    fh = logging.FileHandler(log_file, mode='w', encoding='utf-8')
    fh.setFormatter(logging.Formatter('[%(asctime)s] [%(levelname)s] %(message)s'))
    logger.addHandler(fh)
    
    sh = logging.StreamHandler(sys.stdout)
    sh.setFormatter(logging.Formatter('[%(levelname)s] %(message)s'))
    logger.addHandler(sh)
    
    logger.info("================================================================================")
    logger.info("BM618: GÖREV 1.1 — VERİ DENETİMİ VE DENETLENEBİLİR VERİ TABLOSU HESAPLAMA BAŞLADI")
    logger.info("================================================================================")
    
    all_records = []
    
    for file_year, fname in sorted(RAW_FILES.items()):
        fpath = os.path.join(raw_dir, fname)
        if not os.path.exists(fpath):
            raise FileNotFoundError(f"Ham dosya bulunamadı: {fpath}")
            
        logger.info(f"\nİnceleniyor: {fname} (Dosya Yılı / Kaynak Sürümü: {file_year})")
        wb = openpyxl.load_workbook(fpath, data_only=True)
        sheet_name = wb.sheetnames[0]
        ws = wb[sheet_name]
        
        file_valid_records = []
        file_excluded_rows = []
        
        for r in range(1, ws.max_row + 1):
            row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            # Skip completely empty rows if any
            if not any(v is not None for v in row_vals):
                file_excluded_rows.append((r, row_vals, "Tüm hücreler boş satır"))
                continue
                
            # Column mapping depending on file year
            # 2022 & 2023: No (col 1), İlçe (col 2), Mekan (col 3), Güç (col 4), Ölçü/Birim (col 5), Durum (col 6)
            # 2024 & 2025: İlçe (col 1), Mekan (col 2), Güç (col 3), Ölçü/Birim (col 4), Durum (col 5)
            if file_year in [2022, 2023]:
                d_raw = row_vals[1] if len(row_vals) > 1 else None
                n_raw = row_vals[2] if len(row_vals) > 2 else None
                p_raw = row_vals[3] if len(row_vals) > 3 else None
                s_raw = row_vals[5] if len(row_vals) > 5 else None
            else:
                d_raw = row_vals[0] if len(row_vals) > 0 else None
                n_raw = row_vals[1] if len(row_vals) > 1 else None
                p_raw = row_vals[2] if len(row_vals) > 2 else None
                s_raw = row_vals[4] if len(row_vals) > 4 else None
                
            power_val, is_valid_power = parse_numeric_power(p_raw)
            if not is_valid_power:
                file_excluded_rows.append((r, row_vals, f"Geçersiz/Sayısal olmayan güç değeri ('{p_raw}') - Başlık/Açıklama satırı"))
                continue
                
            # Valid record processing
            d_clean, d_flags = clean_district(d_raw)
            n_clean, n_flags = clean_facility_name(n_raw)
            s_clean, s_flags = clean_status(s_raw)
            
            row_flags = d_flags + n_flags + s_flags
            
            # Check 2024 Çekmeköy Ömerli İtfaiye specific note
            if file_year == 2024 and 'Ömerli İtfaiye' in n_clean:
                row_flags.append("CONFLICT_CANDIDATE_OMERLI_ITFAIYE")
                
            q_flag_str = ";".join(row_flags) if row_flags else "OK"
            
            # Deterministic, traceable record_id
            rec_id = f"REC_{file_year}_R{r:02d}"
            
            record_dict = {
                'record_id': rec_id,
                'source_file': fname,
                'sheet': sheet_name,
                'excel_row': r,
                'file_year': file_year,
                'district_raw': str(d_raw) if d_raw is not None else "",
                'district_clean': d_clean,
                'name_raw': str(n_raw) if n_raw is not None else "",
                'name_clean': n_clean,
                'power_kwp': power_val,
                'status_raw': str(s_raw) if s_raw is not None else "",
                'status_clean': s_clean,
                'entity_group_id': None,  # Görev 1.2'de doldurulacak
                'quality_flag': q_flag_str
            }
            file_valid_records.append(record_dict)
            
        # Log exclusions
        logger.info(f"  Dışlanan başlık/açıklama satırı sayısı: {len(file_excluded_rows)}")
        for r_num, vals, reason in file_excluded_rows:
            logger.info(f"    [Dışlandı] Satır {r_num:02d}: {reason} | İçerik: {vals}")
            
        # Calculate statistics
        df_year = pd.DataFrame(file_valid_records)
        n_rec = len(df_year)
        p_sum = df_year['power_kwp'].sum()
        p_med = df_year['power_kwp'].median()
        p_min = df_year['power_kwp'].min()
        p_max = df_year['power_kwp'].max()
        n_dist = df_year['district_clean'].nunique()
        raw_status_counts = df_year['status_raw'].value_counts().to_dict()
        clean_status_counts = df_year['status_clean'].value_counts().to_dict()
        
        logger.info(f"  Geçerli Tesis Kaydı: {n_rec}")
        logger.info(f"  Benzersiz Temiz İlçe Sayısı: {n_dist}")
        logger.info(f"  Toplam Güç: {p_sum:.1f} kWp (Medyan: {p_med:.1f}, Min: {p_min:.1f}, Maks: {p_max:.1f})")
        logger.info(f"  Ham Durum Dağılımı: {raw_status_counts}")
        logger.info(f"  Temiz Durum Dağılımı: {clean_status_counts}")
        
        # Rigorous assertions against official benchmark values
        exp = EXPECTED_METRICS[file_year]
        assert n_rec == exp['count'], f"{file_year} kayıt sayısı uyuşmuyor: {n_rec} != {exp['count']}"
        assert abs(p_sum - exp['sum']) < 1e-4, f"{file_year} güç toplamı uyuşmuyor: {p_sum} != {exp['sum']}"
        assert n_dist == exp['districts'], f"{file_year} ilçe sayısı uyuşmuyor: {n_dist} != {exp['districts']}"
        assert abs(p_med - exp['median']) < 1e-4, f"{file_year} medyan uyuşmuyor: {p_med} != {exp['median']}"
        assert abs(p_min - exp['min']) < 1e-4, f"{file_year} min uyuşmuyor: {p_min} != {exp['min']}"
        assert abs(p_max - exp['max']) < 1e-4, f"{file_year} maks uyuşmuyor: {p_max} != {exp['max']}"
        for s_k, s_v in exp['status'].items():
            assert raw_status_counts.get(s_k, 0) == s_v, f"{file_year} ham durum '{s_k}' uyuşmuyor"
            
        logger.info(f"  >>> {file_year} YIL VERİLERİ YÖNERGE VE ÖĞRETİM NOTLARIYLA %100 DOĞRULANDI <<<")
        all_records.extend(file_valid_records)

    full_df = pd.DataFrame(all_records)
    
    # Save outputs
    csv_path = os.path.join(output_data_dir, "audited_inventory.csv")
    full_df.to_csv(csv_path, index=False, encoding='utf-8')
    logger.info(f"\nDenetlenebilir envanter tablosu CSV olarak kaydedildi: {csv_path}")
    
    logger.info("================================================================================")
    logger.info(f"GENEL ÖZET: Toplam {len(full_df)} geçerli GES kaydı oluşturuldu.")
    logger.info(f"Toplam Güç: {full_df['power_kwp'].sum():.1f} kWp")
    logger.info("Görev 1.1 Başarıyla Tamamlandı.")
    logger.info("================================================================================")
    
    return full_df


if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    raw_dir = os.path.join(base_dir, "raw")
    data_dir = os.path.join(base_dir, "data")
    logs_dir = os.path.join(base_dir, "logs")
    
    run_audit(raw_dir, data_dir, logs_dir)
