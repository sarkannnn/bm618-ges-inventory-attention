"""
BM618: Az Verili GES Envanterinde Attention, Transformer ve Kanıta Dayalı LLM
Adım 9a — Ham Excel Dosyalarından (raw/) Bağımsız Doğrulama Betiği (verify_raw_excel.py)

Bu betik, data/kb_docs.jsonl veya data/audited_inventory.csv ara tablolarını KULLANMADAN,
doğrudan raw/*.xlsx ham Excel dosyalarını openpyxl ile bağımsız olarak okur ve doğrular:
1. Toplama Soruları: Q06, Q18, Q20, Q22
2. 7 Tekil Test Sorusu: Q11, Q12, Q13, Q14, Q15, Q16, Q24 (Excel satır numarası ve hücre değerleriyle)
3. 3 Çelişki Test Sorusu: Q23, Q25, Q26 (Excel satır numarası ve çelişen değerleriyle)

Ham Dosya Sütun Düzeni:
- 2022 & 2023: Col 1: Sıra No | Col 2: İlçe | Col 3: Tesis Adı | Col 4: Güç | Col 5: Birim | Col 6: Durum
- 2024 & 2025: Col 1: İlçe    | Col 2: Tesis Adı | Col 3: Güç    | Col 4: Birim | Col 5: Durum
"""

import os
import sys
import openpyxl

# UTF-8 konsol çıktısı
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

# Proje kök dizini (mutlak yol stringi olmadan)
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def parse_numeric(val):
    """Excel hücre değerini float'a çevirir (virgül/nokta normalize edilir)."""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return 0.0


def verify_raw_questions():
    print("=" * 80)
    print("HAM EXCEL DOSYALARINDAN (raw/) BAĞIMSIZ DOĞRULAMA (TOPLAMA, TEKİL, ÇELİŞKİ)")
    print("=" * 80)

    raw_dir = os.path.join(PROJECT_ROOT, "raw")
    f_2022 = os.path.join(raw_dir, "2022-yl-kurulumu-tamamlanan-gune-enerji-sistemleri.xlsx")
    f_2023 = os.path.join(raw_dir, "2023_kurulumu-tamamlanan-gune-enerji-sistemleri.xlsx")
    f_2024 = os.path.join(raw_dir, "2024-yl-kurulumu-tamamlanan-gune-enerji-sistemleri.xlsx")
    f_2025 = os.path.join(raw_dir, "2025-yl-kurulumu-tamamlanan-gune-enerji-sistemleri.xlsx")

    wb22 = openpyxl.load_workbook(f_2022, data_only=True)
    ws22 = wb22["2022 KURULUMU TAMAMLANAN GES"]

    wb23 = openpyxl.load_workbook(f_2023, data_only=True)
    ws23 = wb23["2023 KURULUMU TAMAMLANAN GES"]

    wb24 = openpyxl.load_workbook(f_2024, data_only=True)
    ws24 = wb24["2024 KURULUMU TAMAMLANAN GES"]

    wb25 = openpyxl.load_workbook(f_2025, data_only=True)
    ws25 = wb25["2025 KURULUMU TAMAMLANAN GES"]

    # =========================================================================
    # BÖLÜM 1: TOPLAMA SORULARI (Q06, Q18, Q20, Q22)
    # =========================================================================
    print("\n--- BÖLÜM 1: TOPLAMA SORULARI BAĞIMSIZ DOĞRULAMASI ---")

    # Q06: 2022 Dosyası, Fatih İlçe Sütunundaki Kayıtlar
    fatih_rows = []
    fatih_sum = 0.0
    for r in range(2, ws22.max_row + 1):
        dist = ws22.cell(row=r, column=2).value
        name = ws22.cell(row=r, column=3).value
        p_val = ws22.cell(row=r, column=4).value
        if dist and str(dist).strip().lower() == "fatih":
            power = parse_numeric(p_val)
            fatih_rows.append((r, str(name).strip(), power))
            fatih_sum += power
    fatih_sum = round(fatih_sum, 2)
    print(f"\n[Q06] 2022 Dosyası — İlçe Sütununda 'Fatih' Yazan Tesisler:")
    for r_num, name, p in fatih_rows:
        print(f"  - Excel Satır {r_num}: İlçe=Fatih | Mekan={name} | Güç={p:.1f} kWp")
    print(f"  => Ham Excel Toplamı: {fatih_sum:.1f} kWp (Altın Yanıt: 1124.2 kWp)")
    assert abs(fatih_sum - 1124.2) < 1e-4, f"Q06 uyuşmazlığı: {fatih_sum} != 1124.2"
    print("  => [Q06 DOĞRULANDI - BAŞARILI]")

    # Q18: 2024 Dosyası, Beykoz İlçesi
    beykoz_rows = []
    beykoz_sum = 0.0
    for r in range(2, ws24.max_row + 1):
        dist = ws24.cell(row=r, column=1).value
        name = ws24.cell(row=r, column=2).value
        p_val = ws24.cell(row=r, column=3).value
        if dist and str(dist).strip().lower() == "beykoz":
            power = parse_numeric(p_val)
            beykoz_rows.append((r, str(name).strip(), power))
            beykoz_sum += power
    beykoz_sum = round(beykoz_sum, 2)
    print(f"\n[Q18] 2024 Dosyası — Beykoz İlçesi GES Tesisleri:")
    for r_num, name, p in beykoz_rows:
        print(f"  - Excel Satır {r_num}: İlçe=Beykoz | Mekan={name} | Güç={p:.1f} kWp")
    print(f"  => Ham Excel Toplamı: {beykoz_sum:.1f} kWp (Altın Yanıt: 62.0 kWp)")
    assert abs(beykoz_sum - 62.0) < 1e-4, f"Q18 uyuşmazlığı: {beykoz_sum} != 62.0"
    print("  => [Q18 DOĞRULANDI - BAŞARILI]")

    # Q20: 2024 Dosyası, Kartal Yakacık İtfaiye ve Soğanlık Gasilhane
    kartal_rows = []
    kartal_sum = 0.0
    for r in range(2, ws24.max_row + 1):
        dist = ws24.cell(row=r, column=1).value
        name = ws24.cell(row=r, column=2).value
        p_val = ws24.cell(row=r, column=3).value
        if dist and str(dist).strip().lower() == "kartal":
            name_str = str(name)
            if "Yakacık" in name_str or "Soğanlık" in name_str:
                power = parse_numeric(p_val)
                kartal_rows.append((r, str(name).strip(), power))
                kartal_sum += power
    kartal_sum = round(kartal_sum, 2)
    print(f"\n[Q20] 2024 Dosyası — Kartal (Yakacık İtfaiye + Soğanlık Gasilhane):")
    for r_num, name, p in kartal_rows:
        print(f"  - Excel Satır {r_num}: Mekan={name} | Güç={p:.1f} kWp")
    print(f"  => Ham Excel Toplamı: {kartal_sum:.1f} kWp (Altın Yanıt: 62.0 kWp)")
    assert abs(kartal_sum - 62.0) < 1e-4, f"Q20 uyuşmazlığı: {kartal_sum} != 62.0"
    print("  => [Q20 DOĞRULANDI - BAŞARILI]")

    # Q22: 2024 Dosyası, Eyüpsultan Tevfik Aydeniz ve Yeşilpınar Yüzme Havuzu
    eyup_rows = []
    eyup_sum = 0.0
    for r in range(2, ws24.max_row + 1):
        dist = ws24.cell(row=r, column=1).value
        name = ws24.cell(row=r, column=2).value
        p_val = ws24.cell(row=r, column=3).value
        if dist and str(dist).strip().lower() == "eyüpsultan":
            name_str = str(name)
            if "Tevfik Aydeniz" in name_str or "Yeşilpınar" in name_str:
                power = parse_numeric(p_val)
                eyup_rows.append((r, str(name).strip(), power))
                eyup_sum += power
    eyup_sum = round(eyup_sum, 2)
    print(f"\n[Q22] 2024 Dosyası — Eyüpsultan (Tevfik Aydeniz + Yeşilpınar Yüzme Havuzu):")
    for r_num, name, p in eyup_rows:
        print(f"  - Excel Satır {r_num}: Mekan={name} | Güç={p:.1f} kWp")
    print(f"  => Ham Excel Toplamı: {eyup_sum:.1f} kWp (Altın Yanıt: 156.0 kWp)")
    assert abs(eyup_sum - 156.0) < 1e-4, f"Q22 uyuşmazlığı: {eyup_sum} != 156.0"
    print("  => [Q22 DOĞRULANDI - BAŞARILI]")

    # =========================================================================
    # BÖLÜM 2: 7 TEKİL TEST SORUSU (Q11, Q12, Q13, Q14, Q15, Q16, Q24)
    # =========================================================================
    print("\n--- BÖLÜM 2: 7 TEKİL TEST SORUSU HAM EXCEL DOĞRULAMASI ---")

    tekil_specs = [
        ("Q11", ws22, 9, 2, 3, 4, "Sultangazi", "Sultangazi 75. Yıl Spor Salonu GES", 106.0, "2022 dosyasında Sultangazi 75. Yıl Spor Salonu"),
        ("Q12", ws22, 14, 2, 3, 4, "Şile", "Şile Engelliler Yaz Kampı GES", 215.0, "2022 dosyasında Şile Engelliler Yaz Kampı"),
        ("Q13", ws23, 11, 2, 3, 4, "Şişli", "Maçka Deprem Parkı Tuvaleti Çatısı Güneş Enerji Santrali", 6.6, "2023 dosyasında Maçka Deprem Parkı"),
        ("Q14", ws24, 5, 1, 2, 3, "Bahçelievler", "Dr. Enver Ören Kültür Merkezi Çatısı Güneş Enerji Santrali", 180.0, "2024 dosyasında Dr. Enver Ören Kültür Merkezi"),
        ("Q15", ws24, 23, 1, 2, 3, "Kartal", "Bülent Ecevit Kültür Merkezi Çatısı Güneş Enerji Santrali", 149.0, "2024 dosyasında Bülent Ecevit Kültür Merkezi"),
        ("Q16", ws24, 46, 1, 2, 3, "Tuzla", "Bahar Merkezi Çatısı Güneş Enerji Santrali", 458.0, "2024 dosyasında Bahar Merkezi"),
        ("Q24", ws22, 12, 2, 3, 4, "Başakşehir", "Engelliler Merkezi 1. Etap GES", 99.0, "2022 dosyasında Başakşehir Engelliler Merkezi 1. Etap"),
    ]

    for q_id, ws, row_idx, col_d, col_n, col_p, exp_dist, exp_name_part, exp_pwr, desc in tekil_specs:
        act_dist = str(ws.cell(row=row_idx, column=col_d).value).strip()
        act_name = str(ws.cell(row=row_idx, column=col_n).value).strip()
        act_pwr = parse_numeric(ws.cell(row=row_idx, column=col_p).value)

        print(f"\n[{q_id}] {desc}:")
        print(f"  - Excel Dosyası/Sayfası: [{ws.title}], Satır {row_idx}")
        print(f"  - Hücre Değerleri: İlçe='{act_dist}', Mekan='{act_name}', Güç={act_pwr} kWp")
        print(f"  => Beklenen: {exp_pwr} kWp | Okunan: {act_pwr} kWp")

        assert exp_dist.lower() in act_dist.lower(), f"{q_id} ilçe uyuşmazlığı"
        assert exp_name_part.lower() in act_name.lower(), f"{q_id} ad uyuşmazlığı"
        assert abs(act_pwr - exp_pwr) < 1e-4, f"{q_id} güç uyuşmazlığı: {act_pwr} != {exp_pwr}"
        print(f"  => [{q_id} DOĞRULANDI - BAŞARILI]")

    # =========================================================================
    # BÖLÜM 3: 3 ÇELİŞKİ TEST SORUSU (Q23, Q25, Q26)
    # =========================================================================
    print("\n--- BÖLÜM 3: 3 ÇELİŞKİ TEST SORUSU HAM EXCEL DOĞRULAMASI ---")

    # Q23: Florya Atatürk Kent Ormanı (2022 Satır 2 vs 2024 Satır 7)
    q23_p1 = parse_numeric(ws22.cell(row=2, column=4).value)
    q23_n1 = str(ws22.cell(row=2, column=3).value).strip()
    q23_p2 = parse_numeric(ws24.cell(row=7, column=3).value)
    q23_n2 = str(ws24.cell(row=7, column=2).value).strip()
    print(f"\n[Q23] Florya Atatürk Kent Ormanı Yerleşkesi (Ad ve Güç Çelişkisi):")
    print(f"  - 2022 Excel Satır 2: '{q23_n1}' -> {q23_p1} kWp")
    print(f"  - 2024 Excel Satır 7: '{q23_n2}' -> {q23_p2} kWp")
    assert abs(q23_p1 - 20.0) < 1e-4 and abs(q23_p2 - 7.0) < 1e-4
    print("  => [Q23 DOĞRULANDI - BAŞARILI]")

    # Q25: Kartal Yakacık İtfaiye (2024 Satır 24 vs 2025 Satır 42 - Azalan Güç Çelişkisi)
    q25_p1 = parse_numeric(ws24.cell(row=24, column=3).value)
    q25_n1 = str(ws24.cell(row=24, column=2).value).strip()
    q25_p2 = parse_numeric(ws25.cell(row=42, column=3).value)
    q25_n2 = str(ws25.cell(row=42, column=2).value).strip()
    print(f"\n[Q25] Kartal Yakacık İtfaiye (Resmi Kayıtlarda Azalan Güç Çelişkisi):")
    print(f"  - 2024 Excel Satır 24: '{q25_n1}' -> {q25_p1} kWp")
    print(f"  - 2025 Excel Satır 42: '{q25_n2}' -> {q25_p2} kWp")
    assert abs(q25_p1 - 33.0) < 1e-4 and abs(q25_p2 - 29.0) < 1e-4
    print("  => [Q25 DOĞRULANDI - BAŞARILI]")

    # Q26: Çekmeköy Ömerli İtfaiye (2024 Satır 13 vs 2024 Satır 15 - Aynı Tesis-Yıl Güç Çelişkisi)
    q26_p1 = parse_numeric(ws24.cell(row=13, column=3).value)
    q26_n1 = str(ws24.cell(row=13, column=2).value).strip()
    q26_p2 = parse_numeric(ws24.cell(row=15, column=3).value)
    q26_n2 = str(ws24.cell(row=15, column=2).value).strip()
    print(f"\n[Q26] Çekmeköy Ömerli İtfaiye (2024 Dosyası Aynı Tesis-Yıl Güç Çelişkisi):")
    print(f"  - 2024 Excel Satır 13: '{q26_n1}' -> {q26_p1} kWp")
    print(f"  - 2024 Excel Satır 15: '{q26_n2}' -> {q26_p2} kWp")
    assert abs(q26_p1 - 29.0) < 1e-4 and abs(q26_p2 - 25.0) < 1e-4
    print("  => [Q26 DOĞRULANDI - BAŞARILI]")

    print("\n" + "=" * 80)
    print("SONUÇ: TOPLAMA, TEKİL VE ÇELİŞKİ SORULARI HAM EXCEL'DEN BAĞIMSIZ DOĞRULANDI!")
    print("=" * 80)


if __name__ == '__main__':
    verify_raw_questions()
